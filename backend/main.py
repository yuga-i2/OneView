"""AICTE OneView - core backend (FastAPI + SQLite).
Run:  uvicorn main:app --reload     (from the backend folder)
Docs: http://localhost:8000/docs   App: http://localhost:8000
All seeded institutions/cases are SYNTHETIC sample data."""
import os, re, io, csv, json, hmac, time, base64, hashlib, sqlite3
from datetime import datetime, timedelta, timezone
from typing import Optional
from fastapi import FastAPI, HTTPException, Depends, Header, Query
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

HERE = os.path.dirname(os.path.abspath(__file__))
DB = os.environ.get("ONEVIEW_DB", os.path.join(HERE, "oneview.db"))
SECRET = os.environ.get("ONEVIEW_SECRET", "change-me-in-production")
BUREAUS = ["RB", "PGRC", "Approval Bureau", "Academic Bureau"]
CATS = ["Faculty norms", "Fee / admission", "Infrastructure", "Affiliation / approval", "Student grievance", "Fake claims / misinformation"]
STATUSES = ["Pending", "Under Process", "Disposed", "Transferred", "Escalated"]
app = FastAPI(title="AICTE OneView API", version="0.1")

# ---------- db helpers ----------
def conn():
    c = sqlite3.connect(DB); c.row_factory = sqlite3.Row; return c
def q(sql, a=(), one=False):
    with conn() as c:
        rows = [dict(r) for r in c.execute(sql, a).fetchall()]
    return (rows[0] if rows else None) if one else rows
def ex(sql, a=()):
    with conn() as c:
        c.execute(sql, a)
def now(): return datetime.now(timezone.utc).replace(microsecond=0)
def iso(d): return d.isoformat()

# ---------- auth (PBKDF2 passwords, HMAC-signed expiring tokens) ----------
def hpw(pw, salt): return hashlib.pbkdf2_hmac("sha256", pw.encode(), salt.encode(), 100_000).hex()
def make_token(u):
    body = base64.urlsafe_b64encode(json.dumps({"u": u, "e": int(time.time()) + 8 * 3600}).encode()).decode()
    return body + "." + hmac.new(SECRET.encode(), body.encode(), hashlib.sha256).hexdigest()
def current_user(authorization: Optional[str] = Header(None)):
    try:
        body, sig = authorization.split()[1].split(".")
        if not hmac.compare_digest(sig, hmac.new(SECRET.encode(), body.encode(), hashlib.sha256).hexdigest()): raise ValueError
        d = json.loads(base64.urlsafe_b64decode(body))
        if d["e"] < time.time(): raise ValueError
        u = q("SELECT username,name,role,bureau FROM users WHERE username=?", (d["u"],), one=True)
        if u: return u
    except Exception: pass
    raise HTTPException(401, "Please sign in again")
def can_act(u, c): return u["role"] == "admin" or (u["role"] == "officer" and u["bureau"] == c["bureau"])

# ---------- overlap matching: ID (deterministic) -> normalised name (rule-based) ----------
AB = {"inst": "institute", "tech": "technology", "engg": "engineering", "mgmt": "management", "univ": "university"}
SW = {"of", "the", "and", "for", "college"}
def norm(s): return {AB.get(w, w) for w in re.sub(r"[^a-z0-9 ]", " ", s.lower().replace("&", " and ")).split()} - SW
def sim(a, b):
    x, y = norm(a), norm(b); return len(x & y) / (len(x | y) or 1)
def resolve(name, pid):
    if pid and pid.strip():
        r = q("SELECT id FROM institutions WHERE id=?", (pid.strip(),), one=True)
        if r: return r["id"], "AICTE Permanent ID", 1.0
    best = max(q("SELECT id,name FROM institutions"), key=lambda i: sim(name or "", i["name"]), default=None)
    if best and sim(name or "", best["name"]) >= 0.6:
        s = sim(name, best["name"]); return best["id"], f"Name similarity {round(s*100)}%", s
    return None, "", 0.0
def case_out(r):
    d = dict(r); d["age"] = (now() - datetime.fromisoformat(r["received"])).days
    d["overdue"] = r["status"] != "Disposed" and d["age"] > r["sla_days"]
    d["due_in"] = r["sla_days"] - d["age"]; return d
def all_cases():
    return [case_out(r) for r in q("SELECT c.*, i.region, i.name AS inst_name FROM cases c JOIN institutions i ON i.id=c.inst_id ORDER BY received DESC")]
def overlap_groups():
    g = {}
    for c in all_cases():
        if c["status"] != "Disposed": g.setdefault(c["inst_id"], []).append(c)
    dec = {r["inst_id"]: r["decision"] for r in q("SELECT * FROM overlap_decisions")}
    return [{"inst_id": k, "institution": v[0]["inst_name"], "bureaus": sorted({c["bureau"] for c in v}), "decision": dec.get(k),
             "label": "Potential Overlapping Cases - Coordination Required", "cases": v}
            for k, v in g.items() if len({c["bureau"] for c in v}) > 1 and dec.get(k) != "dismissed"]
def log(case_id, user, action, remarks="", ts=None):
    ex("INSERT INTO history(case_id,ts,actor,bureau,action,remarks) VALUES(?,?,?,?,?,?)",
       (case_id, iso(ts or now()), user["name"], user["bureau"] or user["role"], action, remarks))

# ---------- schemas ----------
class Login(BaseModel): username: str; password: str
class NewCase(BaseModel):
    institution_name: str = ""; permanent_id: str = ""; category: str; description: str = ""; bureau: Optional[str] = None
class Action(BaseModel):
    action: str; remarks: str = ""; next_action: Optional[str] = None; to_bureau: Optional[str] = None; status: Optional[str] = None
class Decision(BaseModel): decision: str

# ---------- routes ----------
@app.post("/api/login")
def login(b: Login):
    u = q("SELECT * FROM users WHERE username=?", (b.username,), one=True)
    if not u or not hmac.compare_digest(u["pw"], hpw(b.password, u["salt"])): raise HTTPException(401, "Wrong username or password")
    return {"token": make_token(u["username"]), "user": {k: u[k] for k in ("username", "name", "role", "bureau")}}

@app.get("/api/bootstrap")
def bootstrap(u=Depends(current_user)):
    return {"user": u, "institutions": q("SELECT * FROM institutions"), "cases": all_cases(),
            "decisions": {r["inst_id"]: r["decision"] for r in q("SELECT * FROM overlap_decisions")}}

@app.get("/api/cases")
def list_cases(q_: Optional[str] = Query(None, alias="q"), bureau: Optional[str] = None, region: Optional[str] = None, category: Optional[str] = None,
               status: Optional[str] = None, overdue: Optional[bool] = None, date_from: Optional[str] = None, date_to: Optional[str] = None, u=Depends(current_user)):
    out = []
    for c in all_cases():
        hay = f'{c["id"]} {c["inst_id"]} {c["inst_name"]} {c["name_entered"]}'.lower()
        if q_ and q_.lower() not in hay: continue
        if bureau and c["bureau"] != bureau or region and c["region"] != region or category and c["category"] != category or status and c["status"] != status: continue
        if overdue and not c["overdue"]: continue
        if date_from and c["received"][:10] < date_from or date_to and c["received"][:10] > date_to: continue
        out.append(c)
    return out

@app.get("/api/cases/{cid}")
def get_case(cid: str, u=Depends(current_user)):
    c = next((c for c in all_cases() if c["id"] == cid), None)
    if not c: raise HTTPException(404, "Case not found")
    return {"case": c, "history": q("SELECT * FROM history WHERE case_id=? ORDER BY ts", (cid,)),
            "related": [x for x in all_cases() if x["inst_id"] == c["inst_id"] and x["id"] != cid]}

@app.post("/api/cases")
def create_case(b: NewCase, u=Depends(current_user)):
    if u["role"] not in ("admin", "officer"): raise HTTPException(403, "Your role cannot register complaints")
    if b.category not in CATS: raise HTTPException(422, "Unknown category")
    inst, how, _ = resolve(b.institution_name, b.permanent_id)
    if not inst: raise HTTPException(422, "No matching institution. Check the name or AICTE Permanent ID")
    bureau = b.bureau if u["role"] == "admin" and b.bureau in BUREAUS else u["bureau"]
    n = q("SELECT COUNT(*) n FROM cases", one=True)["n"] + 1
    cid = f"CMP-{now().year}-{n:04d}"
    ex("INSERT INTO cases(id,inst_id,name_entered,has_pid,match_basis,bureau,category,status,received,officer,last_action,next_action,sla_days,description) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
       (cid, inst, b.institution_name or inst, int(bool(b.permanent_id.strip())), how, bureau, b.category, "Pending", iso(now()), u["name"], "Complaint registered", "Assign officer", 30, b.description))
    log(cid, u, "Complaint received", f"Category: {b.category}. Matched to institution by {how}.")
    others = [c for c in all_cases() if c["inst_id"] == inst and c["id"] != cid and c["status"] != "Disposed" and c["bureau"] != bureau]
    for o in others: log(o["id"], u, "Overlap alert", f"New case {cid} registered in {bureau} on the same institution")
    return {"case": next(c for c in all_cases() if c["id"] == cid), "potential_overlaps": others,
            "alert": "Potential Overlapping Cases - Coordination Required" if others else None}

@app.post("/api/cases/{cid}/actions")
def act(cid: str, b: Action, u=Depends(current_user)):
    c = next((c for c in all_cases() if c["id"] == cid), None)
    if not c: raise HTTPException(404, "Case not found")
    if not can_act(u, c): raise HTTPException(403, f"Only {c['bureau']} officers or admins can act on this case")
    st, bur, rem = c["status"], c["bureau"], b.remarks
    if b.action == "transfer":
        if b.to_bureau not in BUREAUS or b.to_bureau == bur: raise HTTPException(422, "Choose a different Bureau/Cell")
        rem = f"Transferred from {bur} to {b.to_bureau}. {rem}"; bur, st = b.to_bureau, "Transferred"
    elif b.action == "escalate": st = "Escalated"
    elif b.action == "dispose": st = "Disposed"
    elif b.action == "update": st = b.status if b.status in STATUSES else ("Under Process" if st in ("Pending", "Transferred") else st)
    else: raise HTTPException(422, "Unknown action")
    ex("UPDATE cases SET status=?,bureau=?,last_action=?,next_action=? WHERE id=?",
       (st, bur, b.action.capitalize() + (": " + b.remarks if b.remarks else ""), b.next_action or ("None" if st == "Disposed" else c["next_action"]), cid))
    log(cid, u, b.action.capitalize(), rem)
    return {"ok": True, "status": st, "bureau": bur}

@app.get("/api/institutions")
def institutions(u=Depends(current_user)):
    cs = all_cases(); out = []
    for i in q("SELECT * FROM institutions"):
        l = [c for c in cs if c["inst_id"] == i["id"]]; o = [c for c in l if c["status"] != "Disposed"]
        out.append({**i, "total": len(l), "open": len(o), "overdue": sum(c["overdue"] for c in l), "bureaus": sorted({c["bureau"] for c in o})})
    return out
@app.get("/api/institutions/{iid}")
def institution(iid: str, u=Depends(current_user)):
    i = q("SELECT * FROM institutions WHERE id=?", (iid,), one=True)
    if not i: raise HTTPException(404, "Institution not found")
    return {**i, "cases": [c for c in all_cases() if c["inst_id"] == iid]}

@app.get("/api/match")
def match(name: str = "", pid: str = "", u=Depends(current_user)):
    inst, how, s = resolve(name, pid)
    if not inst: return {"institution": None}
    return {"institution": q("SELECT * FROM institutions WHERE id=?", (inst,), one=True), "matched_by": how,
            "open_cases": [c for c in all_cases() if c["inst_id"] == inst and c["status"] != "Disposed"]}

@app.get("/api/overlaps")
def overlaps(u=Depends(current_user)): return overlap_groups()
@app.post("/api/overlaps/{iid}")
def decide(iid: str, b: Decision, u=Depends(current_user)):
    if u["role"] not in ("admin", "officer"): raise HTTPException(403, "Your role cannot verify overlaps")
    if b.decision not in ("linked", "dismissed"): raise HTTPException(422, "decision must be linked or dismissed")
    ex("INSERT OR REPLACE INTO overlap_decisions VALUES(?,?,?,?)", (iid, b.decision, u["name"], iso(now())))
    for c in all_cases():
        if c["inst_id"] == iid and c["status"] != "Disposed":
            log(c["id"], u, "Overlap " + b.decision, "Officer verified the potential overlap" if b.decision == "linked" else "Marked as not related")
    return {"ok": True}

@app.get("/api/alerts")
def alerts(u=Depends(current_user)):
    cs = [c for c in all_cases() if c["status"] != "Disposed"]
    return {"overdue": [c for c in cs if c["overdue"]], "due_soon": [c for c in cs if not c["overdue"] and c["due_in"] <= 5],
            "escalated": [c for c in cs if c["status"] == "Escalated"], "overlaps": overlap_groups()}

@app.get("/api/stats")
def stats(u=Depends(current_user)):
    cs = all_cases()
    return {"total": len(cs), **{s: sum(c["status"] == s for c in cs) for s in STATUSES}, "overdue": sum(c["overdue"] for c in cs),
            "overlapping_institutions": len(overlap_groups()), "workload": {b: sum(c["bureau"] == b and c["status"] != "Disposed" for c in cs) for b in BUREAUS}}

@app.get("/api/reports/{dim}")
def report(dim: str, u=Depends(current_user)):
    key = {"bureau": "bureau", "region": "region", "category": "category", "status": "status", "institution": "inst_name"}.get(dim)
    if not key: raise HTTPException(404, "Use bureau, region, category, status or institution")
    g = {}
    for c in all_cases(): g.setdefault(c[key], []).append(c)
    return [{dim: k, "total": len(v), "open": sum(c["status"] != "Disposed" for c in v), "disposed": sum(c["status"] == "Disposed" for c in v),
             "overdue": sum(c["overdue"] for c in v), "disposal_rate_pct": round(100 * sum(c["status"] == "Disposed" for c in v) / len(v))} for k, v in g.items()]

@app.get("/api/export.csv")
def export(u=Depends(current_user)):
    f = io.StringIO(); w = csv.writer(f)
    w.writerow(["Case ID", "Institution", "Permanent ID", "Region", "Bureau", "Category", "Status", "Officer", "Received", "Age (days)", "Overdue"])
    for c in all_cases(): w.writerow([c["id"], c["inst_name"], c["inst_id"], c["region"], c["bureau"], c["category"], c["status"], c["officer"], c["received"][:10], c["age"], c["overdue"]])
    f.seek(0); return StreamingResponse(iter([f.read()]), media_type="text/csv", headers={"Content-Disposition": "attachment; filename=oneview_cases.csv"})

# ---------- Excel / PDF export, periodic reports, document upload ----------
UPLOADS = os.path.join(HERE, "uploads")
ALLOWED = {".pdf", ".png", ".jpg", ".jpeg", ".docx", ".xlsx", ".txt"}
MAX_UPLOAD = 5 * 1024 * 1024
def _case_rows(cs): return [[c["id"], c["inst_name"], c["inst_id"], c["region"], c["bureau"], c["category"], c["status"], c["officer"], c["received"][:10], c["age"], "Yes" if c["overdue"] else "No"] for c in cs]
CASE_HEAD = ["Case ID", "Institution", "Permanent ID", "Region", "Bureau", "Category", "Status", "Officer", "Received", "Age (days)", "Overdue"]
def _period(c, p):
    d = datetime.fromisoformat(c["received"]); return f"{d.year}-{d.month:02d}" if p == "monthly" else f"{d.year}-Q{(d.month - 1) // 3 + 1}"
def _summary(p):
    g = {}
    for c in all_cases(): g.setdefault(_period(c, p), []).append(c)
    return [[k, len(v), sum(c["status"] == "Disposed" for c in v), sum(c["status"] != "Disposed" for c in v), sum(c["overdue"] for c in v)] for k, v in sorted(g.items(), reverse=True)]
SUM_HEAD = ["Period", "Received", "Disposed", "Still open", "Overdue"]

@app.get("/api/export.xlsx")
def export_xlsx(u=Depends(current_user)):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Cases"; ws.append(CASE_HEAD)
    for r in _case_rows(all_cases()): ws.append(r)
    for col, w in zip("ABCDEFGHIJK", (16, 44, 14, 12, 18, 26, 14, 16, 12, 11, 9)): ws.column_dimensions[col].width = w
    f = io.BytesIO(); wb.save(f); f.seek(0)
    return StreamingResponse(f, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=oneview_cases.xlsx"})

@app.get("/api/report-file")
def report_file(period: str = "monthly", fmt: str = "xlsx", u=Depends(current_user)):
    if period not in ("monthly", "quarterly") or fmt not in ("xlsx", "pdf"): raise HTTPException(422, "period: monthly|quarterly, fmt: xlsx|pdf")
    rows, name = _summary(period), f"oneview_{period}_report"
    if fmt == "xlsx":
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.title = period.title(); ws.append(SUM_HEAD)
        for r in rows: ws.append(r)
        ws2 = wb.create_sheet("Cases"); ws2.append(CASE_HEAD)
        for r in _case_rows(all_cases()): ws2.append(r)
        f = io.BytesIO(); wb.save(f); f.seek(0)
        return StreamingResponse(f, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename={name}.xlsx"})
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    f = io.BytesIO(); st = getSampleStyleSheet()
    t = Table([SUM_HEAD] + rows, hAlign="LEFT"); t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0e7a5f")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), .5, colors.HexColor("#d5d9df")), ("ALIGN", (1, 0), (-1, -1), "CENTER")]))
    SimpleDocTemplate(f, pagesize=A4).build([Paragraph(f"AICTE OneView - {period} report", st["Title"]), Paragraph(f"Generated {now().strftime('%d %b %Y')}. Synthetic sample data.", st["Normal"]), Spacer(1, 14), t])
    f.seek(0); return StreamingResponse(f, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={name}.pdf"})

class Doc(BaseModel): filename: str; content_b64: str
@app.get("/api/cases/{cid}/documents")
def list_docs(cid: str, u=Depends(current_user)):
    return q("SELECT id,filename,size,by_user,ts FROM documents WHERE case_id=? ORDER BY id DESC", (cid,))
@app.post("/api/cases/{cid}/documents")
def add_doc(cid: str, b: Doc, u=Depends(current_user)):
    c = next((c for c in all_cases() if c["id"] == cid), None)
    if not c: raise HTTPException(404, "Case not found")
    if not can_act(u, c): raise HTTPException(403, f"Only {c['bureau']} officers or admins can attach documents")
    name = re.sub(r"[^A-Za-z0-9._-]", "_", os.path.basename(b.filename))[:80]
    if os.path.splitext(name)[1].lower() not in ALLOWED: raise HTTPException(422, "Allowed types: " + ", ".join(sorted(ALLOWED)))
    try: data = base64.b64decode(b.content_b64, validate=True)
    except Exception: raise HTTPException(422, "Invalid file data")
    if not data or len(data) > MAX_UPLOAD: raise HTTPException(422, "File must be between 1 byte and 5 MB")
    os.makedirs(UPLOADS, exist_ok=True)
    with conn() as cn:
        i = cn.execute("INSERT INTO documents(case_id,filename,size,by_user,ts) VALUES(?,?,?,?,?)", (cid, name, len(data), u["name"], iso(now()))).lastrowid
    with open(os.path.join(UPLOADS, f"{i}_{name}"), "wb") as fh: fh.write(data)
    log(cid, u, "Document attached", name)
    return {"id": i, "filename": name, "size": len(data)}
@app.get("/api/documents/{did}")
def get_doc(did: int, u=Depends(current_user)):
    d = q("SELECT * FROM documents WHERE id=?", (did,), one=True)
    if not d: raise HTTPException(404, "Document not found")
    return StreamingResponse(open(os.path.join(UPLOADS, f"{d['id']}_{d['filename']}"), "rb"), media_type="application/octet-stream", headers={"Content-Disposition": f"attachment; filename={d['filename']}"})

# ---------- schema + synthetic seed ----------
def init():
    with conn() as c:
        c.execute("CREATE TABLE IF NOT EXISTS documents(id INTEGER PRIMARY KEY AUTOINCREMENT,case_id TEXT,filename TEXT,size INT,by_user TEXT,ts TEXT)")
    with conn() as c:
        c.executescript("""CREATE TABLE IF NOT EXISTS institutions(id TEXT PRIMARY KEY,name TEXT,address TEXT,region TEXT);
        CREATE TABLE IF NOT EXISTS users(username TEXT PRIMARY KEY,name TEXT,role TEXT,bureau TEXT,salt TEXT,pw TEXT);
        CREATE TABLE IF NOT EXISTS cases(id TEXT PRIMARY KEY,inst_id TEXT REFERENCES institutions(id),name_entered TEXT,has_pid INT,match_basis TEXT,bureau TEXT,category TEXT,status TEXT,received TEXT,officer TEXT,last_action TEXT,next_action TEXT,sla_days INT,description TEXT DEFAULT '');
        CREATE TABLE IF NOT EXISTS history(id INTEGER PRIMARY KEY AUTOINCREMENT,case_id TEXT,ts TEXT,actor TEXT,bureau TEXT,action TEXT,remarks TEXT);
        CREATE TABLE IF NOT EXISTS overlap_decisions(inst_id TEXT PRIMARY KEY,decision TEXT,by_user TEXT,ts TEXT);""")
    if q("SELECT 1 FROM institutions LIMIT 1"): return
    I = [("1-4401001", "Sri Venkateswara Institute of Technology", "Sample address, Andhra Pradesh", "Southern"), ("1-5502114", "Maharana Pratap College of Engineering", "Sample address, Rajasthan", "Northern"),
         ("1-7703220", "Deccan Institute of Management & Technology", "Sample address, Maharashtra", "Western"), ("1-3304087", "Bengal School of Engineering and Technology", "Sample address, West Bengal", "Eastern"),
         ("1-6605931", "Narmada Polytechnic College", "Sample address, Madhya Pradesh", "Central"), ("1-4408812", "Kaveri Engineering College", "Sample address, Karnataka", "Southern")]
    for i in I: ex("INSERT INTO institutions VALUES(?,?,?,?)", i)
    for un, nm, role, bu in [("admin", "System Admin", "admin", None), ("rb_officer", "R. Sharma (RB)", "officer", "RB"), ("pgrc_officer", "A. Menon (PGRC)", "officer", "PGRC"),
                             ("approval_officer", "S. Iyer (Approval)", "officer", "Approval Bureau"), ("senior_mgmt", "Dr. K. Verma (Senior Mgmt)", "senior", None), ("regional_south", "N. Reddy (Southern RO)", "regional", None)]:
        ex("INSERT INTO users VALUES(?,?,?,?,?,?)", (un, nm, role, bu, un, hpw("demo123", un)))
    OFF = ["A. Menon", "R. Sharma", "S. Iyer", "P. Das", "K. Verma", "N. Reddy"]; ST = STATUSES
    R = [(1,0,"Sri Venkateswara Institute of Technology",1,0,0,1,34,0,"Notice issued to institution","Await reply",30),(2,0,"Sri Venkateswara Inst. of Tech.",0,1,0,0,9,1,"Complaint registered","Scrutiny of documents",30),
    (3,0,"Sri Venkateswara Institute of Technology",1,2,3,1,12,2,"Approval file requested","Verify affiliation status",45),(4,1,"Maharana Pratap College of Engineering",1,0,1,4,52,3,"Escalated to Member Secretary","Decision on show-cause",30),
    (5,1,"Maharana Pratap College of Engg",0,1,1,1,20,1,"Hearing scheduled","Hearing with complainant",30),(6,2,"Deccan Institute of Management and Technology",1,1,4,2,40,4,"Closed after institution reply","None",30),
    (7,2,"Deccan Inst of Mgmt & Technology",0,0,2,1,6,5,"Inspection ordered","Inspection report",45),(8,3,"Bengal School of Engineering and Technology",1,0,0,0,3,0,"Complaint registered","Assign to officer",30),
    (9,3,"Bengal School of Engineering and Technology",1,3,4,3,25,2,"Transferred to RB","Handled by RB",30),(10,4,"Narmada Polytechnic College",1,1,5,1,47,3,"Clarification sought","Verify website claims",30),
    (11,5,"Kaveri Engineering College",1,0,3,2,60,0,"Disposed with warning","None",45),(12,5,"Kaveri Engineering College",1,1,1,1,8,4,"Reply sought","Review reply",30),
    (13,4,"Narmada Polytechnic College",1,0,0,0,28,5,"Pending allocation","Assign officer",30),(14,2,"Deccan Institute of Management and Technology",1,3,3,1,15,1,"Data validation","Update records",45)]
    for n, k, nm, pid, b, cat, st, age, of, last, nxt, sla in R:
        cid, t0 = f"CMP-2026-{n:04d}", now() - timedelta(days=age)
        ex("INSERT INTO cases(id,inst_id,name_entered,has_pid,match_basis,bureau,category,status,received,officer,last_action,next_action,sla_days) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
           (cid, I[k][0], nm, pid, "AICTE Permanent ID" if pid else "Name similarity", BUREAUS[b], CATS[cat], ST[st], iso(t0), OFF[of], last, nxt, sla))
        sysu = {"name": "System", "bureau": BUREAUS[b], "role": "system"}
        log(cid, sysu, "Complaint received", f"Category: {CATS[cat]}", t0); log(cid, {**sysu, "name": OFF[of]}, "Assigned", "Assigned to " + OFF[of], t0 + timedelta(hours=4))
        if age > 2: log(cid, {**sysu, "name": OFF[of]}, last, "Remarks recorded", t0 + timedelta(days=age // 2))
init()
app.mount("/", StaticFiles(directory=os.path.join(HERE, "..", "frontend"), html=True), name="ui")
